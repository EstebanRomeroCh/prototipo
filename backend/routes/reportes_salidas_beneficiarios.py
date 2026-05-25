# routes/reportes_salidas_beneficiarios.py
#
# Reporte: SALIDAS por BENEFICIARIO (Parroquias/Fundaciones)
# - Por defecto trae TODAS las salidas (desde la primera hasta la última).
# - Incluye información completa del beneficiario (contacto/ubicación/documento).
# - Incluye cuántas familias atiende (detecta automáticamente la columna si existe).
# - Incluye cuánto han entregado por CATEGORÍA (kg y valor) en el mismo endpoint (JSON anidado)
#   y en Excel en una segunda hoja.

from flask import Blueprint, render_template, send_file, session, redirect, url_for, jsonify, request
from database import get_db_connection
import pymysql
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from psycopg2.extras import RealDictCursor


reportes_salidas_benef_bp = Blueprint("reportes_salidas_benef_bp", __name__)


# =========================================================
# Helpers
# =========================================================
def format_kg(value):
    try:
        num = Decimal(str(value or 0))
    except (InvalidOperation, TypeError, ValueError):
        return "0"
    if num == num.to_integral():
        return str(int(num))
    return f"{num:.3f}".rstrip("0").rstrip(".")


def format_money(value):
    try:
        num = Decimal(str(value or 0))
    except (InvalidOperation, TypeError, ValueError):
        return "0"
    text = f"{num:,.0f}"
    return text.replace(",", ".")


def parse_fechas(desde_str: str, hasta_str: str):
    """
    Regla:
      - Si NO envían desde/hasta => (None, None) => SIN filtro por fechas (todas las salidas).
      - Si envían solo una => rango de un día.
      - Si envían ambas => rango normal.
    """
    desde_str = (desde_str or "").strip()
    hasta_str = (hasta_str or "").strip()

    if not desde_str and not hasta_str:
        return None, None

    if desde_str and not hasta_str:
        d = datetime.fromisoformat(desde_str).date()
        return d, d

    if hasta_str and not desde_str:
        h = datetime.fromisoformat(hasta_str).date()
        return h, h

    return (
        datetime.fromisoformat(desde_str).date(),
        datetime.fromisoformat(hasta_str).date(),
    )


def _get_db_name(cursor):
    cursor.execute("SELECT DATABASE() AS db;")
    row = cursor.fetchone() or {}
    return row.get("db")


def _col_exists(cursor, db_name: str, table_name: str, col_name: str) -> bool:
    cursor.execute(
        """
        SELECT 1
        FROM information_schema.COLUMNS
        WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s AND COLUMN_NAME = %s
        LIMIT 1
        """,
        (db_name, table_name, col_name),
    )
    return cursor.fetchone() is not None


def _pick_familias_column(cursor, db_name: str, table_name: str):
    """
    Busca una columna “familias” probable para no depender de un nombre exacto.
    Ajusta esta lista si ya tienes el nombre real.
    """
    candidatos = [
        "familias_atendidas",
        "num_familias",
        "numero_familias",
        "familias",
        "familias_beneficiadas",
        "familias_atendidas_mes",
        "familias_atendidas_total",
    ]
    for c in candidatos:
        if _col_exists(cursor, db_name, table_name, c):
            return c
    return None


def _decimal_or_zero(x):
    try:
        return Decimal(str(x or 0))
    except (InvalidOperation, TypeError, ValueError):
        return Decimal("0")


def cargar_tipos_salida():
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=RealDictCursor)
    try:
        cursor.execute("""
            SELECT id_tipo_salida, nombre
            FROM tipo_salida
            WHERE
                estado IS NULL
                OR LOWER(TRIM(estado)) = 'activo'
                OR TRIM(estado) = '1'
            ORDER BY nombre
        """)
        return cursor.fetchall() or []
    except Exception as e:
        print("Error cargar_tipos_salida:", e)
        return []
    finally:
        cursor.close()
        conn.close()


@reportes_salidas_benef_bp.route("/reportes/salidas/tipos-salida", methods=["GET"])
def api_tipos_salida_reportes():
    if "usuario" not in session:
        return jsonify({"success": False, "message": "Sesión expirada"}), 401

    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=RealDictCursor)
    try:
        cursor.execute("""
            SELECT id_tipo_salida, nombre
            FROM tipo_salida
            WHERE estado IS NULL
                OR LOWER(TRIM(estado))='activo'
                OR TRIM(estado)='1'
            ORDER BY nombre
        """)
        rows = cursor.fetchall() or []
        return jsonify({"success": True, "items": rows})
    except Exception as e:
        print("Error api_tipos_salida_reportes:", e)
        return jsonify({"success": False, "message": "Error al cargar tipos de salida"}), 500
    finally:
        cursor.close()
        conn.close()



# =========================================================
# VISTA HTML
# =========================================================
@reportes_salidas_benef_bp.route("/reportes/salidas/beneficiarios", methods=["GET"])
def vista_reporte_salidas_beneficiarios():
    if "usuario" not in session:
        return redirect(url_for("index"))

    
    return render_template(
        "reporte_salidas_beneficiarios.html",
        usuario=session.get("usuario"),
        rol_nombre=session.get("rol_nombre"),
        tipos_salida=cargar_tipos_salida(),
    )


# =========================================================
# JSON: Beneficiarios + desglose por categoría (kg y valor)
# =========================================================
@reportes_salidas_benef_bp.route("/reportes/salidas/beneficiarios/data", methods=["GET"])
def data_reporte_salidas_beneficiarios():
    if "usuario" not in session:
        return jsonify({"success": False, "message": "Sesión expirada"}), 401

    desde_str = (request.args.get("desde") or "").strip()
    hasta_str = (request.args.get("hasta") or "").strip()

    # filtros opcionales
    q = (request.args.get("q") or "").strip()  # nombre o documento
    estado = (request.args.get("estado") or "").strip()  # Activo|Anulado|""
    id_tipo_salida = request.args.get("id_tipo_salida", type=int)
    beneficiario_tipo = (request.args.get("beneficiario_tipo") or "").strip()  # Parroquia|Fundacion|""

    fecha_desde, fecha_hasta = parse_fechas(desde_str, hasta_str)

    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=RealDictCursor)

    try:
        db_name = _get_db_name(cursor)
        col_p = _pick_familias_column(cursor, db_name, "parroquias") if db_name else None
        col_f = _pick_familias_column(cursor, db_name, "fundacioines") if db_name else None

        # Expresión SQL para “familias atendidas”
        if col_p and col_f:
            familias_expr = f"COALESCE(p.{col_p}, f.{col_f}, 0)"
        elif col_p and not col_f:
            familias_expr = f"COALESCE(p.{col_p}, 0)"
        elif col_f and not col_p:
            familias_expr = f"COALESCE(f.{col_f}, 0)"
        else:
            familias_expr = "0"

        filtros = ["1=1"]
        params = []

        # fechas (solo si se envían)
        if fecha_desde and fecha_hasta:
            filtros.append("DATE(s.fecha_salida) >= %s")
            filtros.append("DATE(s.fecha_salida) <= %s")
            params.extend([fecha_desde, fecha_hasta])

        if estado:
            filtros.append("s.estado = %s")
            params.append(estado)

        if id_tipo_salida:
            filtros.append("s.id_tipo_salida = %s")
            params.append(id_tipo_salida)

        if beneficiario_tipo in ("Parroquia", "Fundacion"):
            if beneficiario_tipo == "Parroquia":
                filtros.append("s.id_parroquia IS NOT NULL")
            else:
                filtros.append("s.id_fundaciones IS NOT NULL")

        if q:
            filtros.append(
                "(COALESCE(p.nombre, f.nombre) LIKE %s OR COALESCE(p.numero_documento, f.numero_documento) LIKE %s)"
            )
            params.extend([f"%{q}%", f"%{q}%"])

        where_clause = "WHERE " + " AND ".join(filtros)

        # Una consulta que trae: beneficiario + categoría (para armar el desglose en el mismo JSON)
        sql = f"""
            SELECT
                -- Beneficiario
                CASE
                    WHEN s.id_parroquia IS NOT NULL THEN 'Parroquia'
                    WHEN s.id_fundaciones IS NOT NULL THEN 'Fundacion'
                    ELSE 'N/A'
                END AS beneficiario_tipo,
                COALESCE(p.id_parroquia, f.id_fundaciones) AS id_beneficiario,
                COALESCE(p.nombre, f.nombre) AS beneficiario_nombre,
                COALESCE(p.nombre_encargado, f.nombre_encargado) AS encargado,
                COALESCE(p.numero_documento, f.numero_documento) AS numero_documento,
                td.nombre AS tipo_documento,
                COALESCE(p.telefono, f.telefono) AS telefono,
                COALESCE(p.direccion, f.direccion) AS direccion,
                COALESCE(p.municipio, f.municipio) AS municipio,
                COALESCE(p.departamento, f.departamento) AS departamento,
                {familias_expr} AS familias_atendidas,

                -- Contexto salida
                COUNT(DISTINCT s.id_salida) AS num_salidas,
                MAX(s.fecha_salida) AS ultima_salida,

                -- Categoría (desglose)
                c.id_categoria,
                c.nombre AS categoria,
                SUM(sd.cantidad_kg) AS total_kg_categoria,
                SUM(COALESCE(sd.valor_total, 0)) AS total_valor_categoria

            FROM salidas s
            LEFT JOIN parroquias p
                ON p.id_parroquia = s.id_parroquia
            LEFT JOIN fundacioines f
                ON f.id_fundaciones = s.id_fundaciones
            LEFT JOIN tipo_documento td
                ON td.id_tipo_doc = COALESCE(p.tipo_doc_id, f.tipo_doc_id)

            JOIN salida_detalles sd
                ON sd.id_salida = s.id_salida
            JOIN categorias c
                ON c.id_categoria = sd.id_categoria

            {where_clause}

            GROUP BY
                beneficiario_tipo, id_beneficiario, beneficiario_nombre, encargado, numero_documento,
                tipo_documento, telefono, direccion, municipio, departamento, familias_atendidas,
                c.id_categoria, c.nombre

            ORDER BY
                beneficiario_tipo ASC,
                beneficiario_nombre ASC,
                c.nombre ASC;
        """

        cursor.execute(sql, params)
        rows = cursor.fetchall() or []

        # Armado de estructura:
        # beneficiarios: [{... , totales, categorias:[...]}]
        benef_map = {}
        tot_kg_general = Decimal("0")
        tot_val_general = Decimal("0")

        for r in rows:
            btipo = r.get("beneficiario_tipo") or "N/A"
            bid = r.get("id_beneficiario")
            if bid is None:
                # seguridad: si hay registros mal formados
                continue

            key = f"{btipo}:{bid}"

            kg_cat = _decimal_or_zero(r.get("total_kg_categoria"))
            val_cat = _decimal_or_zero(r.get("total_valor_categoria"))
            tot_kg_general += kg_cat
            tot_val_general += val_cat

            if key not in benef_map:
                benef_map[key] = {
                    "beneficiario_tipo": btipo,
                    "id_beneficiario": bid,
                    "beneficiario_nombre": r.get("beneficiario_nombre"),
                    "encargado": r.get("encargado"),
                    "tipo_documento": r.get("tipo_documento"),
                    "numero_documento": r.get("numero_documento"),
                    "telefono": r.get("telefono"),
                    "direccion": r.get("direccion"),
                    "municipio": r.get("municipio"),
                    "departamento": r.get("departamento"),
                    "familias_atendidas": int(r.get("familias_atendidas") or 0),
                    "num_salidas": int(r.get("num_salidas") or 0),
                    "ultima_salida": r["ultima_salida"].strftime("%Y-%m-%d") if r.get("ultima_salida") else None,
                    "total_kg": 0.0,
                    "total_valor": 0.0,
                    "total_kg_str": "0",
                    "total_valor_str": "0",
                    "categorias": [],
                }

            benef_map[key]["categorias"].append(
                {
                    "id_categoria": r.get("id_categoria"),
                    "categoria": r.get("categoria"),
                    "total_kg": float(kg_cat),
                    "total_kg_str": format_kg(kg_cat),
                    "total_valor": float(val_cat),
                    "total_valor_str": format_money(val_cat),
                }
            )

            # Acumular totales del beneficiario
            benef_map[key]["total_kg"] += float(kg_cat)
            benef_map[key]["total_valor"] += float(val_cat)

        # Recalcular strings de totales por beneficiario
        beneficiarios = list(benef_map.values())
        for b in beneficiarios:
            b["total_kg_str"] = format_kg(Decimal(str(b["total_kg"])))
            b["total_valor_str"] = format_money(Decimal(str(b["total_valor"])))

        # KPIs: cuántos beneficiarios por tipo
        total_parroquias = sum(1 for b in beneficiarios if b["beneficiario_tipo"] == "Parroquia")
        total_fundaciones = sum(1 for b in beneficiarios if b["beneficiario_tipo"] == "Fundacion")

        return jsonify(
            {
                "success": True,
                "items": beneficiarios,
                "kpis": {
                    "beneficiarios_total": len(beneficiarios),
                    "parroquias_total": total_parroquias,
                    "fundaciones_total": total_fundaciones,
                },
                "totales": {
                    "total_kg": float(tot_kg_general),
                    "total_kg_str": format_kg(tot_kg_general),
                    "total_valor": float(tot_val_general),
                    "total_valor_str": format_money(tot_val_general),
                },
                "meta": {
                    "rango_aplicado": bool(fecha_desde and fecha_hasta),
                    "desde": str(fecha_desde) if fecha_desde else None,
                    "hasta": str(fecha_hasta) if fecha_hasta else None,
                },
            }
        )

    except Exception as e:
        print("Error data_reporte_salidas_beneficiarios:", e)
        return jsonify({"success": False, "message": "Error al generar el reporte"}), 500
    finally:
        cursor.close()
        conn.close()


# =========================================================
# EXCEL: Hoja 1 = Resumen beneficiarios
#        Hoja 2 = Desglose por categoría
# =========================================================
@reportes_salidas_benef_bp.route("/reportes/salidas/beneficiarios/excel", methods=["GET"])
def reporte_salidas_beneficiarios_excel():
    if "usuario" not in session:
        return "Sesión expirada", 401

    desde_str = (request.args.get("desde") or "").strip()
    hasta_str = (request.args.get("hasta") or "").strip()

    q = (request.args.get("q") or "").strip()
    estado = (request.args.get("estado") or "").strip()
    id_tipo_salida = request.args.get("id_tipo_salida", type=int)
    beneficiario_tipo = (request.args.get("beneficiario_tipo") or "").strip()

    fecha_desde, fecha_hasta = parse_fechas(desde_str, hasta_str)

    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=RealDictCursor)

    try:
        db_name = _get_db_name(cursor)
        col_p = _pick_familias_column(cursor, db_name, "parroquias") if db_name else None
        col_f = _pick_familias_column(cursor, db_name, "fundacioines") if db_name else None

        if col_p and col_f:
            familias_expr = f"COALESCE(p.{col_p}, f.{col_f}, 0)"
        elif col_p and not col_f:
            familias_expr = f"COALESCE(p.{col_p}, 0)"
        elif col_f and not col_p:
            familias_expr = f"COALESCE(f.{col_f}, 0)"
        else:
            familias_expr = "0"

        filtros = ["1=1"]
        params = []

        if fecha_desde and fecha_hasta:
            filtros.append("DATE(s.fecha_salida) >= %s")
            filtros.append("DATE(s.fecha_salida) <= %s")
            params.extend([fecha_desde, fecha_hasta])

        if estado:
            filtros.append("s.estado = %s")
            params.append(estado)

        if id_tipo_salida:
            filtros.append("s.id_tipo_salida = %s")
            params.append(id_tipo_salida)

        if beneficiario_tipo in ("Parroquia", "Fundacion"):
            filtros.append("s.id_parroquia IS NOT NULL" if beneficiario_tipo == "Parroquia" else "s.id_fundaciones IS NOT NULL")

        if q:
            filtros.append(
                "(COALESCE(p.nombre, f.nombre) LIKE %s OR COALESCE(p.numero_documento, f.numero_documento) LIKE %s)"
            )
            params.extend([f"%{q}%", f"%{q}%"])

        where_clause = "WHERE " + " AND ".join(filtros)

        # 1) Detalle beneficiario + categoría (para hoja 2)
        sql_det = f"""
            SELECT
                CASE
                    WHEN s.id_parroquia IS NOT NULL THEN 'Parroquia'
                    WHEN s.id_fundaciones IS NOT NULL THEN 'Fundacion'
                    ELSE 'N/A'
                END AS beneficiario_tipo,
                COALESCE(p.id_parroquia, f.id_fundaciones) AS id_beneficiario,
                COALESCE(p.nombre, f.nombre) AS beneficiario_nombre,
                COALESCE(p.nombre_encargado, f.nombre_encargado) AS encargado,
                COALESCE(p.numero_documento, f.numero_documento) AS numero_documento,
                td.nombre AS tipo_documento,
                COALESCE(p.telefono, f.telefono) AS telefono,
                COALESCE(p.direccion, f.direccion) AS direccion,
                COALESCE(p.municipio, f.municipio) AS municipio,
                COALESCE(p.departamento, f.departamento) AS departamento,
                {familias_expr} AS familias_atendidas,
                COUNT(DISTINCT s.id_salida) AS num_salidas,
                MAX(s.fecha_salida) AS ultima_salida,

                c.id_categoria,
                c.nombre AS categoria,
                SUM(sd.cantidad_kg) AS total_kg_categoria,
                SUM(COALESCE(sd.valor_total, 0)) AS total_valor_categoria

            FROM salidas s
            LEFT JOIN parroquias p ON p.id_parroquia = s.id_parroquia
            LEFT JOIN fundacioines f ON f.id_fundaciones = s.id_fundaciones
            LEFT JOIN tipo_documento td
                ON td.id_tipo_doc = COALESCE(p.tipo_doc_id, f.tipo_doc_id)
            JOIN salida_detalles sd ON sd.id_salida = s.id_salida
            JOIN categorias c ON c.id_categoria = sd.id_categoria

            {where_clause}

            GROUP BY
                beneficiario_tipo, id_beneficiario, beneficiario_nombre, encargado, numero_documento,
                tipo_documento, telefono, direccion, municipio, departamento, familias_atendidas,
                c.id_categoria, c.nombre

            ORDER BY
                beneficiario_tipo ASC,
                beneficiario_nombre ASC,
                c.nombre ASC;
        """

        cursor.execute(sql_det, params)
        rows = cursor.fetchall() or []

        # Armar acumulados por beneficiario (para hoja 1)
        benef_map = {}
        for r in rows:
            btipo = r.get("beneficiario_tipo") or "N/A"
            bid = r.get("id_beneficiario")
            if bid is None:
                continue
            key = f"{btipo}:{bid}"

            kg_cat = _decimal_or_zero(r.get("total_kg_categoria"))
            val_cat = _decimal_or_zero(r.get("total_valor_categoria"))

            if key not in benef_map:
                benef_map[key] = {
                    "beneficiario_tipo": btipo,
                    "id_beneficiario": bid,
                    "beneficiario_nombre": r.get("beneficiario_nombre"),
                    "encargado": r.get("encargado"),
                    "tipo_documento": r.get("tipo_documento"),
                    "numero_documento": r.get("numero_documento"),
                    "telefono": r.get("telefono"),
                    "direccion": r.get("direccion"),
                    "municipio": r.get("municipio"),
                    "departamento": r.get("departamento"),
                    "familias_atendidas": int(r.get("familias_atendidas") or 0),
                    "num_salidas": int(r.get("num_salidas") or 0),
                    "ultima_salida": r["ultima_salida"].strftime("%Y-%m-%d") if r.get("ultima_salida") else "",
                    "total_kg": Decimal("0"),
                    "total_valor": Decimal("0"),
                }

            benef_map[key]["total_kg"] += kg_cat
            benef_map[key]["total_valor"] += val_cat

        beneficiarios = list(benef_map.values())

        # Crear Excel
        wb = Workbook()

        # Hoja 1: Resumen por beneficiario
        ws1 = wb.active
        ws1.title = "Resumen beneficiarios"
        ws1.append(
            [
                "Tipo",
                "ID",
                "Beneficiario",
                "Encargado",
                "Tipo doc",
                "Documento",
                "Teléfono",
                "Dirección",
                "Municipio",
                "Departamento",
                "Familias atendidas",
                "N° salidas",
                "Última salida",
                "Total (kg)",
                "Valor total",
            ]
        )

        for b in beneficiarios:
            ws1.append(
                [
                    b["beneficiario_tipo"],
                    b["id_beneficiario"],
                    b["beneficiario_nombre"] or "",
                    b["encargado"] or "",
                    b["tipo_documento"] or "",
                    b["numero_documento"] or "",
                    b["telefono"] or "",
                    b["direccion"] or "",
                    b["municipio"] or "",
                    b["departamento"] or "",
                    int(b["familias_atendidas"] or 0),
                    int(b["num_salidas"] or 0),
                    b["ultima_salida"] or "",
                    float(b["total_kg"]),
                    float(b["total_valor"]),
                ]
            )

        # Gráfica (Top 10) por kg en hoja 1 (si hay datos)
        if beneficiarios:
            # ordenar por kg desc
            beneficiarios_sorted = sorted(beneficiarios, key=lambda x: float(x["total_kg"]), reverse=True)[:10]

            # Escribimos una tablita auxiliar para el chart (para no depender del orden original)
            start_row = len(beneficiarios) + 3
            ws1.cell(row=start_row, column=1, value="Top 10 (kg)")
            ws1.cell(row=start_row + 1, column=1, value="Beneficiario")
            ws1.cell(row=start_row + 1, column=2, value="Kg")

            for i, b in enumerate(beneficiarios_sorted, start=0):
                ws1.cell(row=start_row + 2 + i, column=1, value=(b["beneficiario_nombre"] or "")[:40])
                ws1.cell(row=start_row + 2 + i, column=2, value=float(b["total_kg"]))

            max_row = start_row + 1 + len(beneficiarios_sorted)

            data_ref = Reference(ws1, min_col=2, max_col=2, min_row=start_row + 1, max_row=max_row)
            cats_ref = Reference(ws1, min_col=1, max_col=1, min_row=start_row + 2, max_row=max_row)

            chart = BarChart()
            chart.type = "col"
            chart.title = "Top 10 beneficiarios por kg"
            chart.y_axis.title = "Kg"
            chart.x_axis.title = "Beneficiario"
            chart.add_data(data_ref, titles_from_data=True)
            chart.set_categories(cats_ref)
            chart.height = 12
            chart.width = 26

            ws1.add_chart(chart, "Q2")

        # Hoja 2: Detalle por categoría (kg y valor)
        ws2 = wb.create_sheet("Detalle por categoría")
        ws2.append(
            [
                "Tipo",
                "ID",
                "Beneficiario",
                "Documento",
                "Familias atendidas",
                "Categoría",
                "Total (kg)",
                "Valor total",
            ]
        )

        for r in rows:
            kg_cat = _decimal_or_zero(r.get("total_kg_categoria"))
            val_cat = _decimal_or_zero(r.get("total_valor_categoria"))
            ws2.append(
                [
                    r.get("beneficiario_tipo") or "",
                    r.get("id_beneficiario") or "",
                    r.get("beneficiario_nombre") or "",
                    r.get("numero_documento") or "",
                    int(r.get("familias_atendidas") or 0),
                    r.get("categoria") or "",
                    float(kg_cat),
                    float(val_cat),
                ]
            )

        # Guardar a memoria
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        # nombre
        if fecha_desde and fecha_hasta:
            filename = f"reporte_salidas_beneficiarios_{fecha_desde.strftime('%Y%m%d')}_{fecha_hasta.strftime('%Y%m%d')}.xlsx"
        else:
            filename = "reporte_salidas_beneficiarios_TODO.xlsx"

        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    except Exception as e:
        print("Error reporte_salidas_beneficiarios_excel:", e)
        return "Error al generar el Excel", 500
    finally:
        cursor.close()
        conn.close()
